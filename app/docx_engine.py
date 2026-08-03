from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.field_utils import dropdown_option_values
from app.word_control_utils import (
    classify_native_control,
    get_control_identifier,
    iter_unique_story_roots,
    read_dropdown_options,
)


PLACEHOLDER_PATTERN = re.compile(
    r"\{\{([^{}]+)\}\}"
)
REPEAT_MARKER_PATTERN = re.compile(
    r"\{\{\s*repeat:([A-Za-z][A-Za-z0-9_.-]*)\s*\}\}",
    re.IGNORECASE,
)
ROW_NUMBER_IDS = {
    "row.number",
    "row.index",
    "linha.numero",
    "linha.número",
}


class DocumentGenerationError(RuntimeError):
    """
    Raised when a document cannot be generated.
    """


def generate_docx(
    template_path: Path,
    output_path: Path,
    values: dict[str, Any],
) -> None:
    """
    Fill a DOCX template and save a completed copy.

    XML-level paragraph replacement is used so placeholders in text boxes,
    shapes, headers, footers, and tables are also replaced.
    """

    template_path = Path(template_path)
    output_path = Path(output_path)

    if not template_path.exists():
        raise DocumentGenerationError(
            f"O modelo não existe: {template_path}"
        )

    try:
        document = Document(str(template_path))

        for root in iter_unique_story_roots(document):
            _expand_repeatable_rows(
                root,
                values,
            )

            _replace_placeholders_in_root(
                root,
                values,
            )

            _replace_native_controls(
                root,
                values,
            )

            _replace_legacy_checkbox_controls(
                root,
                values,
            )

        unresolved = _find_unresolved_placeholders(
            document
        )

        if unresolved:
            raise DocumentGenerationError(
                "Ainda existem campos não resolvidos: "
                + ", ".join(sorted(unresolved))
            )

        _set_document_text_color(
            document,
            "000000",
        )

        output_path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        document.save(str(output_path))

    except DocumentGenerationError:
        raise

    except Exception as exc:
        raise DocumentGenerationError(
            f"Não foi possível gerar o documento: {exc}"
        ) from exc


def _expand_repeatable_rows(
    root,
    values: dict[str, Any],
) -> None:
    """Duplicate Word table rows marked with ``{{repeat:field.id}}``.

    The complete ``w:tr`` element is copied, preserving borders, widths,
    shading, paragraph formatting, and merged header rows. Only the marked
    model row is repeated; headers remain untouched.
    """

    template_rows = [
        row
        for row in root.iter(qn("w:tr"))
        if REPEAT_MARKER_PATTERN.search(
            _visible_element_text(row)
        )
    ]

    for template_row in template_rows:
        row_text = _visible_element_text(template_row)
        repeat_matches = list(
            REPEAT_MARKER_PATTERN.finditer(row_text)
        )
        table_ids = {
            match.group(1).strip()
            for match in repeat_matches
        }
        if len(table_ids) != 1:
            raise DocumentGenerationError(
                "Uma linha repetível deve usar apenas um marcador repeat."
            )

        table_id = next(iter(table_ids))
        if table_id not in values:
            raise DocumentGenerationError(
                f"Nenhuma linha foi informada para a tabela repetível '{table_id}'."
            )

        raw_rows = values.get(table_id)
        if not isinstance(raw_rows, list):
            raise DocumentGenerationError(
                f"A tabela repetível '{table_id}' recebeu dados em formato inválido."
            )

        parent = template_row.getparent()
        if parent is None:
            continue
        insert_at = parent.index(template_row)

        for row_index, row_values in enumerate(raw_rows, start=1):
            if not isinstance(row_values, dict):
                raise DocumentGenerationError(
                    f"O item {row_index} da tabela '{table_id}' não é uma linha válida."
                )

            cloned_row = deepcopy(template_row)
            scoped_values = dict(values)
            for column_id, value in row_values.items():
                scoped_values[
                    f"{table_id}.{str(column_id).strip()}"
                ] = value

            number = str(
                row_values.get("__row_number__")
                or f"{row_index:02d}"
            )
            for number_id in ROW_NUMBER_IDS:
                scoped_values[number_id] = number

            _replace_placeholders_in_root(
                cloned_row,
                scoped_values,
            )
            _replace_native_controls(
                cloned_row,
                scoped_values,
            )
            _replace_legacy_checkbox_controls(
                cloned_row,
                scoped_values,
            )

            unresolved = _unresolved_in_element(
                cloned_row
            )
            if unresolved:
                raise DocumentGenerationError(
                    f"A linha {row_index} da tabela '{table_id}' possui "
                    "marcadores sem conteúdo: "
                    + ", ".join(sorted(unresolved))
                )

            parent.insert(insert_at, cloned_row)
            insert_at += 1

        parent.remove(template_row)


def _visible_element_text(element) -> str:
    return "".join(
        node.text or ""
        for node in element.iter()
        if node.tag in {
            qn("w:t"),
            qn("w:instrText"),
        }
    )


def _unresolved_in_element(element) -> set[str]:
    unresolved: set[str] = set()
    for paragraph in element.iter(qn("w:p")):
        text = "".join(
            node.text or ""
            for node in _paragraph_text_elements(paragraph)
        )
        for match in PLACEHOLDER_PATTERN.finditer(text):
            unresolved.add(match.group(1).strip())
    return unresolved


def _replace_placeholders_in_root(
    root,
    values: dict[str, Any],
) -> None:
    """
    Replace placeholders in every paragraph, including paragraphs located
    inside text boxes.
    """

    for paragraph_element in root.iter(qn("w:p")):
        _replace_in_paragraph_element(
            paragraph_element,
            values,
        )


def _paragraph_text_elements(
    paragraph_element,
) -> list[Any]:
    """
    Return text nodes belonging to this paragraph only.

    Nested w:p elements, such as text-box paragraphs inside a drawing, are
    excluded and processed independently.
    """

    text_elements: list[Any] = []

    def walk(element) -> None:
        for child in element.iterchildren():
            if child.tag == qn("w:p"):
                continue

            if child.tag in {
                qn("w:t"),
                qn("w:instrText"),
            }:
                text_elements.append(child)
                continue

            walk(child)

    walk(paragraph_element)
    return text_elements


def _replace_in_paragraph_element(
    paragraph_element,
    values: dict[str, Any],
) -> None:
    text_elements = _paragraph_text_elements(
        paragraph_element
    )

    if not text_elements:
        return

    original_text = "".join(
        element.text or ""
        for element in text_elements
    )

    updated_text = PLACEHOLDER_PATTERN.sub(
        lambda match: _replace_placeholder(
            match,
            values,
        ),
        original_text,
    )

    if updated_text == original_text:
        return

    _set_text_with_breaks(
        text_elements[0],
        updated_text,
    )

    for element in text_elements[1:]:
        element.text = ""


def _replace_placeholder(
    match,
    values: dict[str, Any],
) -> str:
    raw_value = match.group(1).strip()

    if raw_value.casefold().startswith("repeat:"):
        return ""

    if raw_value.lower().startswith("checkbox:"):
        field_id = raw_value.split(":", 1)[1].strip()

        if field_id not in values:
            return match.group(0)

        return (
            "☑"
            if _value_as_bool(values[field_id])
            else "☐"
        )

    if raw_value.lower().startswith("date:"):
        field_id = raw_value.split(":", 1)[1].strip()

        if field_id not in values:
            return match.group(0)

        return str(values[field_id] or "")

    if raw_value.lower().startswith("dropdown:"):
        definition = raw_value.split(":", 1)[1]

        parts = [
            part.strip()
            for part in definition.split("|")
        ]

        field_id = parts[0] if parts else ""

        options = dropdown_option_values(parts[1:])

        if field_id not in values:
            return match.group(0)

        selected_value = _validated_selected_text(
            values, field_id, options
        )

        return selected_value

    field_id = raw_value

    if field_id not in values:
        return match.group(0)

    value = values[field_id]

    if isinstance(value, bool):
        return "☑" if value else "☐"

    if value is None:
        return ""

    return str(value)


def _validate_dropdown_selection(
    field_id: str,
    selected_value: str,
    options: list[str],
) -> None:
    if selected_value and options and selected_value not in options:
        raise DocumentGenerationError(
            f"A opção '{selected_value}' não é válida "
            f"para a lista suspensa '{field_id}'."
        )


def _validated_selected_text(
    values: dict[str, Any],
    field_id: str,
    options: list[str],
) -> str:
    selected_value = str(values[field_id] or "").strip()
    _validate_dropdown_selection(
        field_id,
        selected_value,
        options,
    )
    return selected_value


def _replace_native_controls(
    root,
    values: dict[str, Any],
) -> None:
    for sdt_element in root.iter(qn("w:sdt")):
        properties = sdt_element.find(qn("w:sdtPr"))

        if properties is None:
            continue

        control_type, control_element = classify_native_control(
            properties
        )
        if control_type is None:
            continue

        field_id = get_control_identifier(
            sdt_element
        )

        if not field_id:
            raise DocumentGenerationError(
                'Um campo nativo do Word não possui Marca nem Título.'
            )

        if field_id not in values:
            raise DocumentGenerationError(
                f"Nenhum conteúdo foi informado para '{field_id}'."
            )

        if control_type == "checkbox":
            _set_native_checkbox_state(
                sdt_element,
                _value_as_bool(
                    values[field_id]
                ),
            )
            continue

        if control_type == "date":
            _set_content_control_text(
                sdt_element,
                str(values[field_id] or ""),
            )
            continue

        options = read_dropdown_options(
            control_element
        )

        selected_value = _validated_selected_text(
            values, field_id, options
        )

        _set_content_control_text(
            sdt_element,
            selected_value,
        )



def _value_as_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    return str(value).strip().casefold() in {
        "1",
        "true",
        "yes",
        "sim",
        "checked",
        "marcado",
        "x",
        "☑",
    }


def _set_native_checkbox_state(
    sdt_element,
    checked: bool,
) -> None:
    properties = sdt_element.find(qn("w:sdtPr"))
    if properties is None:
        return

    checkbox = properties.find(qn("w14:checkbox"))
    if checkbox is None:
        return

    checked_element = checkbox.find(qn("w14:checked"))
    if checked_element is None:
        checked_element = OxmlElement("w14:checked")
        checkbox.insert(0, checked_element)

    checked_element.set(
        qn("w14:val"),
        "1" if checked else "0",
    )
    _set_content_control_text(
        sdt_element,
        "☑" if checked else "☐",
    )


def _replace_legacy_checkbox_controls(
    root,
    values: dict[str, Any],
) -> None:
    """
    Update legacy Word FORMCHECKBOX fields without removing any boxes.
    """

    for fld_char in root.iter(qn("w:fldChar")):
        ff_data = fld_char.find(qn("w:ffData"))

        if ff_data is None:
            continue

        checkbox = ff_data.find(qn("w:checkBox"))

        if checkbox is None:
            continue

        name = ff_data.find(qn("w:name"))

        if name is None:
            continue

        field_id = _normalize_checkbox_identifier(
            name.get(qn("w:val"), "")
        )

        if not field_id:
            continue

        checked = _value_as_bool(
            values.get(field_id, False)
        )

        checked_element = checkbox.find(
            qn("w:checked")
        )

        if checked_element is None:
            checked_element = OxmlElement(
                "w:checked"
            )
            checkbox.append(checked_element)

        checked_element.set(
            qn("w:val"),
            "1" if checked else "0",
        )

        default_element = checkbox.find(
            qn("w:default")
        )

        if default_element is None:
            default_element = OxmlElement(
                "w:default"
            )
            checkbox.append(default_element)

        default_element.set(
            qn("w:val"),
            "1" if checked else "0",
        )


def _normalize_checkbox_identifier(
    value: str,
) -> str:
    value = str(value).strip()

    if value.lower().startswith("checkbox:"):
        return value.split(":", 1)[1].strip()

    return value

def _set_content_control_text(
    sdt_element,
    value: str,
) -> None:
    properties = sdt_element.find(qn("w:sdtPr"))

    if properties is not None:
        showing_placeholder = properties.find(
            qn("w:showingPlcHdr")
        )

        if showing_placeholder is not None:
            properties.remove(showing_placeholder)

    content = sdt_element.find(
        qn("w:sdtContent")
    )

    if content is None:
        return

    text_elements = list(
        content.iter(qn("w:t"))
    )

    if text_elements:
        _set_text_with_breaks(
            text_elements[0],
            value,
        )

        for text_element in text_elements[1:]:
            text_element.text = ""

        return

    paragraph = next(
        content.iter(qn("w:p")),
        None,
    )

    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    run.append(text)
    _set_text_with_breaks(text, value)

    if paragraph is not None:
        paragraph.append(run)
    else:
        content.append(run)


def _set_text_with_breaks(
    text_element,
    value: Any,
) -> None:
    """Write text and convert embedded newlines to Word line breaks."""
    normalized = str(value or "").replace("\r\n", "\n").replace("\r", "\n")
    segments = normalized.split("\n")
    parent = text_element.getparent()

    first = segments[0] if segments else ""
    text_element.text = first
    if first.startswith(" ") or first.endswith(" "):
        text_element.set(
            "{http://www.w3.org/XML/1998/namespace}space",
            "preserve",
        )

    if parent is None or len(segments) <= 1:
        return

    insert_at = parent.index(text_element) + 1
    for segment in segments[1:]:
        line_break = OxmlElement("w:br")
        parent.insert(insert_at, line_break)
        insert_at += 1

        next_text = OxmlElement("w:t")
        next_text.text = segment
        if segment.startswith(" ") or segment.endswith(" "):
            next_text.set(
                "{http://www.w3.org/XML/1998/namespace}space",
                "preserve",
            )
        parent.insert(insert_at, next_text)
        insert_at += 1


def _set_document_text_color(
    document,
    color_hex: str = "000000",
) -> None:
    """Apply one explicit text color to every Word run in the output.

    Template authors frequently use red or another highlight color while
    placing markers.  Generated documents should not inherit those editing
    colors, so all text runs in the main document, tables, headers, footers,
    text boxes, and other story parts are normalized to black while keeping
    font family, size, emphasis, borders, and paragraph formatting intact.
    """

    normalized_color = str(color_hex).strip().lstrip("#").upper()
    if not re.fullmatch(r"[0-9A-F]{6}", normalized_color):
        raise ValueError(
            "A cor do texto deve usar seis dígitos hexadecimais."
        )

    for root in iter_unique_story_roots(document):
        for run_element in root.iter(qn("w:r")):
            run_properties = run_element.find(qn("w:rPr"))
            if run_properties is None:
                run_properties = OxmlElement("w:rPr")
                run_element.insert(0, run_properties)

            color_element = run_properties.find(qn("w:color"))
            if color_element is None:
                color_element = OxmlElement("w:color")
                run_properties.append(color_element)

            color_element.set(
                qn("w:val"),
                normalized_color,
            )

            for attribute_name in (
                "w:themeColor",
                "w:themeTint",
                "w:themeShade",
            ):
                color_element.attrib.pop(
                    qn(attribute_name),
                    None,
                )


def _find_unresolved_placeholders(
    document,
) -> set[str]:
    unresolved: set[str] = set()

    for root in iter_unique_story_roots(document):
        for paragraph_element in root.iter(qn("w:p")):
            text = "".join(
                element.text or ""
                for element in _paragraph_text_elements(
                    paragraph_element
                )
            )

            for match in PLACEHOLDER_PATTERN.finditer(text):
                unresolved.add(match.group(1))

    return unresolved
